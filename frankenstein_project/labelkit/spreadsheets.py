"""Spreadsheet upload parsing."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any, Dict, List

from fastapi import HTTPException


def _read_spreadsheet_bytes(filename: str, data: bytes) -> List[Dict[str, Any]]:
    suffix = Path(filename or "").suffix.lower()
    rows: List[List[Any]] = []

    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [row for row in reader]
    elif suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook  # type: ignore
        except Exception as exc:
            raise HTTPException(status_code=500, detail="Excel upload requires openpyxl in the runtime.") from exc

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    else:
        raise HTTPException(status_code=400, detail="Upload an .xlsx, .xlsm, or .csv file.")

    while rows and not any(str(cell or "").strip() for cell in rows[0]):
        rows.pop(0)
    if not rows:
        return []

    headers = [str(cell or "").strip() for cell in rows[0]]
    out: List[Dict[str, Any]] = []
    for raw_row in rows[1:]:
        row = {
            headers[index]: raw_row[index]
            for index in range(min(len(headers), len(raw_row)))
            if headers[index]
        }
        if any(str(value or "").strip() for value in row.values()):
            out.append(row)
    return out
